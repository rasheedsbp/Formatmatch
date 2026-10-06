"""Write a prepared Document as a LaTeX project (main.tex + figures + refs.bib) in a zip."""
from __future__ import annotations

import io
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from typing import Dict, List

from . import model as M
from .classify import UNNUMBERED, known_section
from .model import Run
from .refs import _expand

NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

UNI = {
    "“": "``", "”": "''", "‘": "`", "’": "'", "–": "--", "—": "---", "…": "\\ldots{}", "\u00a0": "~",
    "¹": "\\textsuperscript{1}", "²": "\\textsuperscript{2}", "³": "\\textsuperscript{3}",
    "⁴": "\\textsuperscript{4}", "⁵": "\\textsuperscript{5}", "⁶": "\\textsuperscript{6}",
    "±": "$\\pm$", "×": "$\\times$", "≤": "$\\leq$", "≥": "$\\geq$", "≈": "$\\approx$", "→": "$\\rightarrow$",
    "←": "$\\leftarrow$", "∼": "$\\sim$", "°": "$^\\circ$", "µ": "$\\mu$", "−": "$-$", "·": "$\\cdot$",
    "α": "$\\alpha$", "β": "$\\beta$", "γ": "$\\gamma$", "δ": "$\\delta$", "ε": "$\\epsilon$", "θ": "$\\theta$",
    "λ": "$\\lambda$", "μ": "$\\mu$", "π": "$\\pi$", "σ": "$\\sigma$", "τ": "$\\tau$", "φ": "$\\phi$",
    "ω": "$\\omega$", "Δ": "$\\Delta$", "Σ": "$\\Sigma$", "Ω": "$\\Omega$", "∈": "$\\in$", "∑": "$\\sum$",
    "√": "$\\surd$", "∞": "$\\infty$", "•": "\\textbullet{}", "§": "\\S{}", "©": "\\textcopyright{}",
    "€": "\\texteuro{}", "₹": "Rs.", "™": "\\texttrademark{}", "®": "\\textregistered{}",
}
SPECIAL = {"\\": "\\textbackslash{}", "&": "\\&", "%": "\\%", "$": "\\$", "#": "\\#", "_": "\\_",
           "{": "\\{", "}": "\\}", "~": "\\textasciitilde{}", "^": "\\textasciicircum{}"}


def esc(s: str) -> str:
    out = []
    for ch in s:
        if ch in SPECIAL:
            out.append(SPECIAL[ch])
        elif ch in UNI:
            out.append(UNI[ch])
        elif ch == "\t":
            out.append(" ")
        else:
            out.append(ch)
    return "".join(out)


# ------------------------------------------------------------------ OMML → LaTeX (batched through pandoc)
def omml_to_latex(omml_list: List[str]) -> Dict[int, str]:
    if not omml_list or not shutil.which("pandoc"):
        return {}
    import docx
    from docx.oxml import parse_xml
    d = docx.Document()
    for k, x in enumerate(omml_list):
        d.add_paragraph(f"FMEQSTART{k}FMEQ")
        p = d.add_paragraph()
        p._p.append(parse_xml(x))
    with tempfile.TemporaryDirectory() as td:
        src = os.path.join(td, "eq.docx")
        d.save(src)
        r = subprocess.run(["pandoc", src, "-t", "latex", "--wrap=none"], capture_output=True, text=True, timeout=120)
        if r.returncode != 0:
            return {}
        tex = r.stdout
    out = {}
    parts = re.split(r"FMEQSTART(\d+)FMEQ", tex)
    for i in range(1, len(parts), 2):
        k = int(parts[i])
        body = parts[i + 1].strip()
        m = re.search(r"\\\[(.*?)\\\]", body, re.S) or re.search(r"\\\((.*?)\\\)", body, re.S) or \
            re.search(r"\$\$(.*?)\$\$", body, re.S) or re.search(r"\$(.*?)\$", body, re.S)
        if m:
            out[k] = m.group(1).strip()
    return out


class LatexRenderer:
    def __init__(self, spec: dict, refs=None):
        self.spec = spec
        self.refs = refs or []
        self.files: Dict[str, bytes] = {}
        self.math: Dict[int, str] = {}
        self.math_ids: Dict[int, int] = {}
        self.cls = spec["latex"].get("class", "article")

    # ---------------------------------------------------------- inline text
    def runs(self, runs: List[Run]) -> str:
        out = []
        for r in runs:
            if r.is_math:
                lx = r.latex or self.math.get(self.math_ids.get(id(r), -1))
                out.append(f"\\({lx}\\)" if lx else "\\(\\square\\)")
                continue
            if r.cite and self.mode_numeric:
                nums = re.findall(r"\d+(?:\s*[–-]\s*\d+)?", r.text)
                keys = []
                for n in nums:
                    keys += [f"r{x}" for x in _expand(n.replace("–", "-"))]
                if keys:
                    out.append(("\\textsuperscript{" if r.sup else "") + "\\cite{" + ",".join(keys) + "}" +
                               ("}" if r.sup else ""))
                    continue
            t = esc(r.text)
            if not t:
                continue
            if r.sup:
                t = f"\\textsuperscript{{{t}}}"
            if r.sub:
                t = f"\\textsubscript{{{t}}}"
            if r.italic:
                t = f"\\emph{{{t}}}"
            if r.bold:
                t = f"\\textbf{{{t}}}"
            out.append(t)
        s = "".join(out)
        s = re.sub(r"\\textbf\{([^{}]*)\}\\textbf\{", r"\\textbf{\1", s)
        return s.strip()

    @property
    def mode_numeric(self):
        return self.spec["citations"]["mode"] in ("numeric", "superscript")

    # ---------------------------------------------------------- preamble/front matter
    def preamble(self, front: dict) -> str:
        S = self.spec
        cls = self.cls
        opts = S["latex"].get("options", "")
        pg = S["page"]
        L = [f"\\documentclass[{opts}]{{{cls}}}" if opts else f"\\documentclass{{{cls}}}",
             "% Generated by FormatMatch — CS Software Solutions (cssoftwaresolutions.in)"]
        pk = ["\\usepackage[utf8]{inputenc}", "\\usepackage[T1]{fontenc}", "\\usepackage{graphicx}",
              "\\usepackage{amsmath,amssymb}", "\\usepackage{booktabs}", "\\usepackage{textcomp}",
              "\\usepackage{url}"]
        if cls not in ("IEEEtran", "llncs", "sn-jnl", "elsarticle"):
            fam = S["font"].get("family", "")
            if "Times" in fam:
                pk.append("\\IfFileExists{newtxtext.sty}{\\usepackage{newtxtext,newtxmath}}{\\usepackage{mathptmx}}")
            elif "Palatino" in fam:
                pk.append("\\IfFileExists{newpxtext.sty}{\\usepackage{newpxtext,newpxmath}}{\\usepackage{mathpazo}}")
            elif fam in ("Arial", "Helvetica"):
                pk.append("\\usepackage{helvet}\\renewcommand{\\familydefault}{\\sfdefault}")
            m = pg.get("margins_cm", [2.54] * 4)
            paper = "a4paper" if pg.get("size") == "A4" else "letterpaper"
            pk.append(f"\\usepackage[{paper},top={m[0]}cm,bottom={m[1]}cm,left={m[2]}cm,right={m[3]}cm,"
                      f"columnsep={pg.get('col_gap_cm', 0.5)}cm]{{geometry}}")
            ls = S["font"].get("line_spacing", 1.0)
            pk.append("\\usepackage{setspace}" + ("\\doublespacing" if ls >= 2 else "\\onehalfspacing" if ls >= 1.5 else ""))
            if pg.get("line_numbers"):
                pk.append("\\usepackage{lineno}\\linenumbers")
            if S["headings"].get("scheme") == "none":
                pk.append("\\setcounter{secnumdepth}{0}")
            pk.append("\\usepackage{caption}")
            cap = S["captions"]
            pk.append(f"\\captionsetup[figure]{{name={{{esc(cap['figure_label'].rstrip('.'))}}}" +
                      (",labelsep=period" if cap.get("separator", ". ").strip() == "." else "") +
                      (",labelfont=bf" if cap.get("label_bold") else "") + "}")
            pk.append(f"\\captionsetup[table]{{name={{{esc(cap['table_label'])}}}" +
                      (",labelfont=bf" if cap.get("label_bold") else "") + "}")
            if cap.get("figure_label", "").startswith("Fig."):
                pk.append("\\captionsetup[figure]{name=Fig.}")
        if cls == "elsarticle":
            pk.append("\\usepackage{lineno}")
        if cls == "sn-jnl":
            pk = [x for x in pk if "inputenc" not in x and "fontenc" not in x]
        L += pk
        if cls == "IEEEtran":
            L.append("\\hyphenation{op-tical net-works semi-conduc-tor}")
        L.append("\\begin{document}")
        return "\n".join(L)

    def front_matter(self, front: dict) -> str:
        cls = self.cls
        T = self.runs(front["title"].runs) if front.get("title") else "Title"
        authors = [self.runs(b.runs) for b in front.get("authors", [])]
        affils = [self.runs(b.runs) for b in front.get("affils", [])]
        abstract = front.get("abstract", [])
        abs_tex = "\n\n".join((f"\\textbf{{{esc(b.label)}}} -- " if b.label else "") + self.runs(b.runs) for b in abstract)
        kws = front.get("keywords")
        kw_list = [k.strip() for k in re.split(r"\s*[;,·]\s*", kws.text.rstrip(".")) if k.strip()] if kws else []
        L = []
        if cls == "IEEEtran":
            L.append(f"\\title{{{T}}}")
            auth = " \\\\\n".join(authors + [f"\\textit{{{a}}}" for a in affils])
            L.append(f"\\author{{{auth}}}")
            L.append("\\maketitle")
            if abs_tex:
                L.append(f"\\begin{{abstract}}\n{abs_tex}\n\\end{{abstract}}")
            if kw_list:
                L.append("\\begin{IEEEkeywords}\n" + ", ".join(esc(k) for k in kw_list) + "\n\\end{IEEEkeywords}")
        elif cls == "elsarticle":
            L.append("\\begin{frontmatter}")
            L.append(f"\\title{{{T}}}")
            for a in authors:
                for name in re.split(r",\s*(?:and\s+)?|\s+and\s+", a):
                    if name.strip():
                        L.append(f"\\author{{{name.strip()}}}")
            for a in affils:
                L.append(f"\\affiliation{{organization={{{a}}}}}")
            if abs_tex:
                L.append(f"\\begin{{abstract}}\n{abs_tex}\n\\end{{abstract}}")
            if kw_list:
                L.append("\\begin{keyword}\n" + " \\sep ".join(esc(k) for k in kw_list) + "\n\\end{keyword}")
            L.append("\\end{frontmatter}")
            if self.spec["page"].get("line_numbers"):
                L.append("\\linenumbers")
        elif cls == "llncs":
            L.append(f"\\title{{{T}}}")
            L.append("\\author{" + " \\and ".join(authors) + "}")
            L.append("\\institute{" + " \\\\ ".join(affils) + "}")
            L.append("\\maketitle")
            if abs_tex:
                L.append("\\begin{abstract}\n" + abs_tex + "\n" +
                         ("\\keywords{" + " \\and ".join(esc(k) for k in kw_list) + "}\n" if kw_list else "") +
                         "\\end{abstract}")
        elif cls == "sn-jnl":
            L.append(f"\\title{{{T}}}")
            for a in authors:
                L.append(f"\\author{{{a}}}")
            for k, a in enumerate(affils, 1):
                L.append(f"\\affil[{k}]{{{a}}}")
            if abs_tex:
                L.append(f"\\abstract{{{abs_tex}}}")
            if kw_list:
                L.append("\\keywords{" + ", ".join(esc(k) for k in kw_list) + "}")
            L.append("\\maketitle")
        else:
            L.append(f"\\title{{{T}}}")
            L.append("\\author{" + " \\\\\n".join(authors + [f"\\small\\textit{{{a}}}" for a in affils]) + "}")
            L.append("\\date{}")
            L.append("\\maketitle")
            if abs_tex:
                ab = self.spec["abstract"]
                if ab.get("layout") == "inline":
                    L.append(f"\\noindent\\textbf{{\\textit{{{esc(ab.get('heading', 'Abstract'))}{esc(ab.get('suffix', '—'))}}}}}{abs_tex}\n")
                else:
                    L.append(f"\\begin{{abstract}}\n{abs_tex}\n\\end{{abstract}}")
            if kw_list:
                kwspec = self.spec["keywords"]
                L.append(f"\\noindent\\textbf{{{esc(kwspec.get('label', 'Keywords'))}{esc(kwspec.get('suffix', ': '))}}}" +
                         esc(kwspec.get("separator", "; ")).join(esc(k) for k in kw_list) + "\n")
        return "\n".join(L)

    # ---------------------------------------------------------- body
    def body(self, blocks: List[M.Block]) -> str:
        L = []
        fig_n = tab_n = 0
        i = 0
        cmds = ["section", "subsection", "subsubsection", "paragraph"]
        two_col = int(self.spec["page"].get("columns", 1)) > 1 or self.cls == "IEEEtran"
        while i < len(blocks):
            b = blocks[i]
            r = b.role
            if r == M.HEADING:
                if b.text.strip().lower() in ("references", "bibliography", "reference") or \
                        (i + 1 < len(blocks) and blocks[i + 1].role == M.REFERENCE):
                    i += 1
                    continue
                cmd = cmds[min(b.level, 4) - 1]
                star = "*" if (not b.prefix and (known_section(b.text) in UNNUMBERED or self.spec["headings"]["scheme"] == "none")) else ""
                txt = self.runs(b.runs)
                if self.cls == "IEEEtran" and b.level == 1:
                    txt = txt.title() if txt.isupper() else txt
                L.append(f"\n\\{cmd}{star}{{{txt}}}")
            elif r == M.PARA:
                L.append("\n" + self.runs(b.runs))
            elif r == M.LIST:
                env = "enumerate" if b.ordered else "itemize"
                L.append(f"\\begin{{{env}}}")
                while i < len(blocks) and blocks[i].role == M.LIST:
                    L.append("  \\item " + self.runs(blocks[i].runs))
                    i += 1
                L.append(f"\\end{{{env}}}")
                continue
            elif r == M.EQUATION:
                body = " ".join((self.math.get(self.math_ids.get(id(x), -1)) or x.latex or "") for x in b.runs if x.is_math)
                if b.eq_num:
                    L.append(f"\\begin{{equation}}\n{body}\n\\label{{eq:{b.eq_num}}}\n\\end{{equation}}")
                else:
                    L.append(f"\\[\n{body}\n\\]")
            elif r in (M.FIGURE, M.CAPTION, M.TABLE):
                # gather a float group: caption + float(s) in any order
                grp = []
                while i < len(blocks) and blocks[i].role in (M.FIGURE, M.CAPTION, M.TABLE):
                    grp.append(blocks[i]); i += 1
                    kinds = {x.role for x in grp}
                    if M.CAPTION in kinds and (M.FIGURE in kinds or M.TABLE in kinds):
                        # stop once caption+float paired (unless more images of the same figure follow)
                        if not (i < len(blocks) and blocks[i].role == M.FIGURE and M.FIGURE in kinds and grp[-1].role == M.FIGURE):
                            break
                L.extend(self.float_group(grp, two_col))
                continue
            i += 1
        return "\n".join(L)

    def float_group(self, grp, two_col):
        cap = next((x for x in grp if x.role == M.CAPTION), None)
        figs = [x for x in grp if x.role == M.FIGURE]
        tabs = [x for x in grp if x.role == M.TABLE]
        out = []
        cap_tex = self.runs(cap.runs) if cap else ""
        if figs:
            out.append("\\begin{figure}[!t]\n\\centering")
            w = 1.0 / len(figs) - 0.02 if len(figs) > 1 else 1.0
            for f in figs:
                name = self.add_image(f)
                if name:
                    out.append(f"\\includegraphics[width={w:.2f}\\linewidth]{{{name}}}")
                else:
                    out.append("\\fbox{Figure could not be exported}")
            if cap is not None:
                out.append(f"\\caption{{{cap_tex}}}\\label{{fig:{cap.cap_num}}}")
            out.append("\\end{figure}")
        for t in tabs:
            ncol = max(len(r) for r in t.rows) if t.rows else 1
            env = "table*" if (two_col and ncol > 5) else "table"
            out.append(f"\\begin{{{env}}}[!t]\n\\centering")
            if cap is not None:
                out.append(f"\\caption{{{cap_tex}}}\\label{{tab:{cap.cap_num}}}")
            colspec = "l" + "c" * (ncol - 1)
            body = [f"\\begin{{tabular}}{{{colspec}}}", "\\toprule"]
            for k, row in enumerate(t.rows):
                cells = [esc(c) for c in row] + [""] * (ncol - len(row))
                if k == 0:
                    cells = [f"\\textbf{{{c}}}" if c else c for c in cells]
                body.append(" & ".join(cells) + " \\\\")
                if k == 0:
                    body.append("\\midrule")
            body += ["\\bottomrule", "\\end{tabular}"]
            if ncol > 4:
                out.append("\\resizebox{\\linewidth}{!}{%\n" + "\n".join(body) + "}")
            else:
                out.extend(body)
            out.append(f"\\end{{{env}}}")
        if cap is not None and not figs and not tabs:
            out.append(f"% Orphan caption: {cap_tex}")
        return out

    def add_image(self, f: M.Block):
        ext = (f.image_ext or "png").lower()
        data = f.image
        if ext not in ("png", "jpg", "jpeg", "pdf"):
            try:
                from PIL import Image
                im = Image.open(io.BytesIO(data))
                buf = io.BytesIO(); im.convert("RGB").save(buf, "PNG"); data = buf.getvalue(); ext = "png"
            except Exception:
                return None
        n = len([k for k in self.files if k.startswith("figures/")]) + 1
        name = f"figures/fig{n}.{ext}"
        self.files[name] = data
        return name

    # ---------------------------------------------------------- references
    def bibliography(self, refs_blocks: List[M.Block]) -> str:
        if not refs_blocks:
            return ""
        L = [f"\\begin{{thebibliography}}{{{len(refs_blocks)}}}"]
        if self.cls not in ("IEEEtran",):
            head = self.spec["references"].get("heading", "References")
            if self.cls == "article":
                L.insert(0, f"\\renewcommand{{\\refname}}{{{esc(head)}}}")
        for k, b in enumerate(refs_blocks, 1):
            runs = list(b.runs)
            if runs and re.match(r"^\s*(\[\d+\]|\d+\.)\s*", runs[0].text):
                runs = [Run(re.sub(r"^\s*(\[\d+\]|\d+\.)\s*", "", runs[0].text), runs[0].bold, runs[0].italic)] + runs[1:]
            L.append(f"\\bibitem{{r{k}}} " + self.runs(runs))
        L.append("\\end{thebibliography}")
        return "\n".join(L)

    def bibtex(self) -> str:
        out = []
        for k, r in enumerate(self.refs, 1):
            typ = {"conference": "inproceedings", "book": "book"}.get(r.kind, "article")
            f = {"author": " and ".join(f"{fam}, {ini}" if ini else f"{{{fam}}}" for fam, ini in r.authors),
                 "title": "{" + r.title + "}", "year": r.year}
            if typ == "article":
                f["journal"] = r.container
            elif typ == "inproceedings":
                f["booktitle"] = r.container
            for key in ("volume", "pages", "doi"):
                v = getattr(r, key)
                if v:
                    f[key if key != "issue" else "number"] = v.replace("–", "--")
            if r.issue:
                f["number"] = r.issue
            fields = ",\n".join(f"  {k2} = {{{v}}}" for k2, v in f.items() if v)
            out.append(f"@{typ}{{r{k},\n{fields}\n}}")
        return "\n\n".join(out)

    # ---------------------------------------------------------- main
    def render(self, doc: M.Document) -> bytes:
        maths = []
        for b in doc.blocks:
            for r in b.runs:
                if r.omml:
                    self.math_ids[id(r)] = len(maths)
                    maths.append(r.omml)
        self.math = omml_to_latex(maths)
        front = {"title": doc.first(M.TITLE), "authors": doc.by_role(M.AUTHOR), "affils": doc.by_role(M.AFFIL),
                 "abstract": doc.by_role(M.ABSTRACT), "keywords": doc.first(M.KEYWORDS)}
        front_roles = {M.TITLE, M.AUTHOR, M.AFFIL, M.ABSTRACT, M.KEYWORDS, M.OTHER_FRONT}
        body_blocks = [b for b in doc.blocks if b.role not in front_roles and b.role != M.REFERENCE]
        refs = doc.by_role(M.REFERENCE)
        tex = [self.preamble(front), self.front_matter(front), self.body(body_blocks), "", self.bibliography(refs),
               "", "\\end{document}", ""]
        main = "\n".join(tex)
        self.files["main.tex"] = main.encode("utf-8")
        if self.refs:
            bst = self.spec["latex"].get("bibstyle") or "plain"
            self.files["refs.bib"] = self.bibtex().encode("utf-8")
        readme = (
            "FormatMatch LaTeX export\n========================\n\n"
            f"Document class: {self.cls} ({self.spec['latex'].get('options', '')})\n"
            "Compile with: pdflatex main.tex (twice). On Overleaf, upload this folder as a new project.\n\n"
            "The reference list is written as \\begin{thebibliography} so it already matches the target style.\n"
            f"refs.bib is included if you prefer BibTeX: replace the thebibliography block with\n"
            f"  \\bibliographystyle{{{self.spec['latex'].get('bibstyle') or 'plain'}}}\n  \\bibliography{{refs}}\n"
            "and keep the \\cite{rN} keys.\n\n"
            "If the class file (e.g. IEEEtran.cls, elsarticle.cls, llncs.cls, sn-jnl.cls) is not installed, download\n"
            "the publisher's LaTeX template and copy its .cls/.bst files into this folder.\n"
        )
        self.files["README.txt"] = readme.encode()
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for name, data in self.files.items():
                z.writestr(name, data)
        return buf.getvalue()


def render_latex(doc: M.Document, spec: dict, refs=None) -> bytes:
    return LatexRenderer(spec, refs).render(doc)


def try_compile(zip_bytes: bytes, fallback_class: bool = True):
    """Compile check (used in tests / optional preview). Returns (ok, pdf_bytes|None, log_tail)."""
    if not shutil.which("pdflatex"):
        return False, None, "pdflatex not installed"
    with tempfile.TemporaryDirectory() as td:
        with zipfile.ZipFile(io.BytesIO(zip_bytes)) as z:
            z.extractall(td)
        main = os.path.join(td, "main.tex")
        src = open(main, encoding="utf-8").read()
        m = re.search(r"\\documentclass(\[[^\]]*\])?\{([^}]+)\}", src)
        cls = m.group(2) if m else "article"
        r = subprocess.run(["kpsewhich", cls + ".cls"], capture_output=True, text=True)
        if not r.stdout.strip() and not os.path.exists(os.path.join(td, cls + ".cls")) and fallback_class:
            src = _degrade_to_article(src, cls)
            open(main, "w", encoding="utf-8").write(src)
        log = ""
        for _ in range(2):
            p = subprocess.run(["pdflatex", "-interaction=nonstopmode", "-halt-on-error", "main.tex"], cwd=td,
                               capture_output=True, text=True, timeout=180)
            log = p.stdout[-3000:]
            if p.returncode != 0:
                return False, None, log
        pdf = os.path.join(td, "main.pdf")
        return os.path.exists(pdf), open(pdf, "rb").read() if os.path.exists(pdf) else None, log


def _degrade_to_article(src: str, cls: str) -> str:
    """Make a class-specific file compile with plain article (for previews only)."""
    src = re.sub(r"\\documentclass(\[[^\]]*\])?\{[^}]+\}", r"\\documentclass[10pt]{article}", src, count=1)
    src = src.replace("\\begin{IEEEkeywords}", "\\par\\noindent\\textit{Index Terms}---").replace("\\end{IEEEkeywords}", "")
    src = re.sub(r"\\begin\{frontmatter\}|\\end\{frontmatter\}", "", src)
    src = re.sub(r"\\affiliation\{organization=\{(.*?)\}\}", r"\\date{\1}", src)
    src = re.sub(r"\\begin\{keyword\}(.*?)\\end\{keyword\}", lambda m: "\\par\\noindent\\textit{Keywords:} " + m.group(1).replace("\\sep", ";"), src, flags=re.S)
    src = re.sub(r"\\author\{([^{}]*)\}\s*\n\\author\{", r"\\author{\1 \\and ", src)
    src = src.replace("\\institute{", "\\date{").replace("\\keywords{", "\\par\\noindent\\textit{Keywords:} {")
    src = re.sub(r"\\affil\[\d+\]\{(.*?)\}", r"\\date{\1}", src)
    src = src.replace("\\hyphenation{op-tical net-works semi-conduc-tor}", "")
    src = re.sub(r"\\abstract\{(.*?)\}\n", lambda m: "\\begin{abstract}" + m.group(1) + "\\end{abstract}\n", src, flags=re.S)
    return src
