"""DOCX → Document model.

Equations (OMML, MathType / Equation Editor OLE objects, equation images), equation tables, algorithm /
pseudocode blocks and complex tables are kept *verbatim* (their source XML is carried to the output), so the
formatter never reflows or restyles them.
"""
from __future__ import annotations

import collections
import copy
import re
from typing import List, Optional

import docx
from docx.oxml.ns import qn
from lxml import etree

from . import model as M
from .algo import MONO_RE, is_algo_head, is_algo_line
from .classify import CAPTION_RE, ParaInfo, RoleClassifier, strip_heading_number
from .runs_util import merge_adjacent, strip_trailing, trim_leading

NS_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
NS_V = "urn:schemas-microsoft-com:vml"
NS_O = "urn:schemas-microsoft-com:office:office"
EQ_NUM_RE = re.compile(r"^\s*[\(\[]\s*(\d+[a-z]?|[A-Z]?\d+(?:[.\-]\d+)*[a-z]?)\s*[\)\]]\s*$")
EQ_PROGID_RE = re.compile(r"equation|mathtype|dsmt|mathml|ee\b", re.I)
SMALL_IMG_IN = 0.9   # images shorter than this (inches) inside text are treated as inline maths/symbols


def _style_chain_attr(style, getter):
    s = style
    while s is not None:
        v = getter(s)
        if v is not None:
            return v
        s = s.base_style
    return None


def _vml_size(el):
    """(w_in, h_in) from a VML shape style string."""
    for shp in el.iter(f"{{{NS_V}}}shape"):
        st = shp.get("style", "")
        def g(k):
            m = re.search(k + r"\s*:\s*([\d.]+)\s*(pt|in|px)?", st)
            if not m:
                return None
            v = float(m.group(1)); u = m.group(2) or "pt"
            return v / 72 if u == "pt" else (v if u == "in" else v / 96)
        return g("width"), g("height")
    return None, None


def has_math(el) -> bool:
    if next(el.iter(f"{{{NS_M}}}oMath"), None) is not None:
        return True
    for ole in el.iter(f"{{{NS_O}}}OLEObject"):
        if EQ_PROGID_RE.search(ole.get("ProgID", "")):
            return True
    return False


class DocxReader:
    def __init__(self, path_or_stream):
        self.d = docx.Document(path_or_stream)
        self.part = self.d.part
        self.default_size = self._default_size()
        self.numfmt_cache = {}

    # ---------------- helpers ----------------
    def _default_size(self) -> float:
        try:
            st = self.d.styles["Normal"]
            sz = _style_chain_attr(st, lambda s: s.font.size)
            if sz:
                return sz.pt
        except KeyError:
            pass
        rpr = self.d.styles.element.find(qn("w:docDefaults") + "/" + qn("w:rPrDefault") + "/" + qn("w:rPr"))
        if rpr is not None:
            sz = rpr.find(qn("w:sz"))
            if sz is not None:
                return int(sz.get(qn("w:val"))) / 2
        return 11.0

    def _para_style(self, p):
        try:
            return p.style
        except Exception:
            return None

    def _style_size(self, style) -> float:
        if style is None:
            return self.default_size
        sz = _style_chain_attr(style, lambda s: s.font.size)
        return sz.pt if sz else self.default_size

    def _style_bold(self, style) -> bool:
        if style is None:
            return False
        return bool(_style_chain_attr(style, lambda s: s.font.bold))

    def _is_ordered(self, numId: str, ilvl: str) -> bool:
        key = (numId, ilvl)
        if key in self.numfmt_cache:
            return self.numfmt_cache[key]
        ordered = False
        try:
            numbering = self.part.numbering_part.element
            num = numbering.find(f'{qn("w:num")}[@{qn("w:numId")}="{numId}"]')
            if num is not None:
                absid = num.find(qn("w:abstractNumId")).get(qn("w:val"))
                absn = numbering.find(f'{qn("w:abstractNum")}[@{qn("w:abstractNumId")}="{absid}"]')
                lvl = absn.find(f'{qn("w:lvl")}[@{qn("w:ilvl")}="{ilvl}"]') if absn is not None else None
                if lvl is not None:
                    fmt = lvl.find(qn("w:numFmt"))
                    ordered = fmt is not None and fmt.get(qn("w:val")) not in ("bullet", "none")
        except Exception:
            pass
        self.numfmt_cache[key] = ordered
        return ordered

    # ---------------- run extraction ----------------
    def _runs(self, p_el, para_bold: bool, keep_breaks: bool = False):
        """Runs of a paragraph. Images/OLE objects become placeholder runs (raw=<w:r>, img=info)."""
        runs: List[M.Run] = []

        def rpr_flags(r_el):
            rpr = r_el.find(qn("w:rPr"))
            b = para_bold
            i = sup = sub = False
            if rpr is not None:
                be = rpr.find(qn("w:b"))
                if be is not None:
                    b = be.get(qn("w:val")) not in ("0", "false")
                ie = rpr.find(qn("w:i"))
                if ie is not None:
                    i = ie.get(qn("w:val")) not in ("0", "false")
                va = rpr.find(qn("w:vertAlign"))
                if va is not None:
                    sup = va.get(qn("w:val")) == "superscript"
                    sub = va.get(qn("w:val")) == "subscript"
            return b, i, sup, sub

        def walk(el):
            for ch in el:
                if not isinstance(ch.tag, str):
                    continue
                q = etree.QName(ch)
                tag, ns = q.localname, q.namespace
                if ns == NS_M and tag in ("oMath", "oMathPara"):
                    runs.append(M.Run(omml=etree.tostring(ch, encoding="unicode")))
                elif tag == "r":
                    b, i, sup, sub = rpr_flags(ch)
                    buf = []
                    for t in ch:
                        tl = etree.QName(t).localname if isinstance(t.tag, str) else ""
                        if tl == "t":
                            buf.append(t.text or "")
                        elif tl == "tab":
                            buf.append("\t" if keep_breaks else " ")
                        elif tl in ("br", "cr"):
                            buf.append("\n" if keep_breaks else " ")
                        elif tl == "noBreakHyphen":
                            buf.append("-")
                        elif tl == "sym":
                            buf.append(chr(int(t.get(qn("w:char"), "20"), 16)) if t.get(qn("w:char")) else "")
                        elif tl in ("drawing", "pict", "object"):
                            if buf:
                                runs.append(M.Run("".join(buf), b, i, sup, sub)); buf = []
                            info = self._object_info(t)
                            if info is not None:
                                runs.append(M.Run(raw=copy.deepcopy(ch), img=info))
                    if buf:
                        runs.append(M.Run("".join(buf), b, i, sup, sub))
                elif tag in ("hyperlink", "smartTag", "ins", "fldSimple", "customXml", "sdt", "sdtContent", "bdo", "dir"):
                    walk(ch)
                elif tag == "AlternateContent":
                    for sub in ch:
                        if etree.QName(sub).localname == "Choice":
                            walk(sub)
                            break
        walk(p_el)
        return merge_adjacent(runs)

    def _object_info(self, el) -> Optional[dict]:
        """Describe an embedded drawing / VML picture / OLE object."""
        kind = "pic"
        ole = next(el.iter(f"{{{NS_O}}}OLEObject"), None)
        if ole is not None:
            kind = "ole-eq" if EQ_PROGID_RE.search(ole.get("ProgID", "")) else "ole"
        blob, ext, w, h = None, "png", None, None
        for blip in el.iter(f"{{{NS_A}}}blip"):
            rid = blip.get(f"{{{NS_R}}}embed") or blip.get(f"{{{NS_R}}}link")
            if rid and rid in self.part.related_parts:
                ip = self.part.related_parts[rid]
                blob, ext = ip.blob, ip.partname.ext.lower()
            break
        for ex in el.iter(f"{{{NS_WP}}}extent"):
            try:
                w, h = int(ex.get("cx")) / 914400, int(ex.get("cy")) / 914400
            except Exception:
                pass
            break
        if blob is None:
            for imd in el.iter(f"{{{NS_V}}}imagedata"):
                rid = imd.get(f"{{{NS_R}}}id")
                if rid and rid in self.part.related_parts:
                    ip = self.part.related_parts[rid]
                    blob, ext = ip.blob, ip.partname.ext.lower()
                break
            if w is None:
                w, h = _vml_size(el)
        if blob is None and kind == "pic":
            return None     # e.g. text box / shape without picture
        inline = next(el.iter(f"{{{NS_WP}}}anchor"), None) is None
        return {"kind": kind, "blob": blob, "ext": ext, "w": w, "h": h, "inline": inline}

    # ---------------- paragraph facts (for algorithm detection) ----------------
    def _pfacts(self, p_el):
        from docx.text.paragraph import Paragraph
        p = Paragraph(p_el, self.d._body)
        st = self._para_style(p)
        style_name = st.name if st is not None else ""
        runs = self._runs(p_el, self._style_bold(st))
        text = "".join(r.text for r in runs if not r.is_math)
        ppr = p_el.find(qn("w:pPr"))
        ind = 0.0
        numbered = False
        if ppr is not None:
            ie = ppr.find(qn("w:ind"))
            if ie is not None:
                for k in ("w:left", "w:start"):
                    v = ie.get(qn(k))
                    if v and v.lstrip("-").isdigit():
                        ind = max(ind, int(v) / 1440)
            numbered = ppr.find(qn("w:numPr")) is not None
        lead_tabs = len(re.match(r"^[\t ]*", "".join(r.text for r in runs if not r.is_math)).group(0).replace("    ", "\t").replace(" ", ""))
        fonts = " ".join(rf.get(qn("w:ascii"), "") for rf in p_el.iter(qn("w:rFonts")))
        nonempty = [r for r in runs if r.text.strip() and not r.is_math]
        bold = bool(nonempty) and all(r.bold for r in nonempty)
        return {"text": text.strip(), "style": style_name, "indented": ind > 0.05 or lead_tabs > 0,
                "numbered": numbered and not style_name.lower().startswith("heading"),
                "math": has_math(p_el), "mono": bool(MONO_RE.search(fonts)), "bold": bold,
                "breaks": len(list(p_el.iter(qn("w:br"))))}

    def _algo_lines(self, elements) -> list:
        """[(indent_level, runs)] for the paragraphs of an algorithm (keeps line breaks)."""
        out = []
        for p_el in elements:
            ps = [p_el] if etree.QName(p_el).localname == "p" else list(p_el.iter(qn("w:p")))
            for pe in ps:
                ppr = pe.find(qn("w:pPr"))
                base = 0.0
                if ppr is not None and ppr.find(qn("w:ind")) is not None:
                    v = ppr.find(qn("w:ind")).get(qn("w:left")) or ppr.find(qn("w:ind")).get(qn("w:start"))
                    if v and v.lstrip("-").isdigit():
                        base = max(0, int(v)) / 360       # 0.25 in per level
                runs = self._runs(pe, False, keep_breaks=True)
                line: List[M.Run] = []
                for r in runs:
                    if r.is_math or "\n" not in r.text:
                        line.append(r); continue
                    parts = r.text.split("\n")
                    for k, part in enumerate(parts):
                        if k > 0:
                            out.append(self._finish_line(line, base)); line = []
                        if part:
                            nr = copy.copy(r); nr.text = part; line.append(nr)
                out.append(self._finish_line(line, base))
        return [l for l in out if l[1]]

    @staticmethod
    def _finish_line(runs, base):
        lead = 0
        if runs and not runs[0].is_math:
            m = re.match(r"^[\t ]*", runs[0].text)
            lead = m.group(0).count("\t") + m.group(0).count(" ") // 4
            runs = [copy.copy(runs[0])] + runs[1:]
            runs[0].text = runs[0].text.lstrip("\t ")
        for r in runs:
            if not r.is_math:
                r.text = r.text.replace("\t", " ")
        runs = [r for r in runs if r.text or r.is_math]
        return (round(base + lead, 2), runs)

    # ---------------- main ----------------
    def read(self) -> M.Document:
        doc = M.Document(source_type="docx", src_docx=self.d)
        body = self.d.element.body
        items = []
        for el in body.iterchildren():
            tag = etree.QName(el).localname
            if tag in ("p", "tbl"):
                items.append(el)
            elif tag == "sdt":
                content = el.find(qn("w:sdtContent"))
                if content is not None:
                    items += [s for s in content.iterchildren() if etree.QName(s).localname in ("p", "tbl")]

        # body size = dominant run size weighted by text length
        sizes = collections.Counter()
        for p in self.d.paragraphs:
            st = self._para_style(p)
            base = self._style_size(st)
            for r in p.runs:
                sz = r.font.size.pt if r.font.size else base
                sizes[round(sz * 2) / 2] += len(r.text)
        body_size = sizes.most_common(1)[0][0] if sizes else self.default_size
        max_size = max(sizes) if sizes else body_size
        clf = RoleClassifier(body_size=body_size, max_size=max_size)

        from docx.text.paragraph import Paragraph

        # expand layout tables (figure boxes) into their paragraphs; classify the others
        stream = []
        for el in items:
            if etree.QName(el).localname == "p":
                stream.append(("p", el))
                continue
            kind = self._table_kind(el)
            if kind == "layout":
                stream += [("p", cp) for cp in el.iter(qn("w:p"))]
            else:
                stream.append((kind, el))

        i = 0
        n = len(stream)
        while i < n:
            kind, el = stream[i]
            if kind == "algo_tbl":
                self._emit_algorithm(doc, [el])
                i += 1
                continue
            if kind == "eq_tbl":
                self._emit_eq_table(doc, el)
                i += 1
                continue
            if kind == "tbl":
                self._emit_table(doc, el)
                i += 1
                continue
            # paragraph: algorithm head?
            if clf.state not in ("front", "abstract", "refs"):
                f = self._pfacts(el)
                if is_algo_head(f["text"], f["bold"]):
                    j = i + 1
                    group = [el]
                    while j < n and stream[j][0] == "p":
                        g = self._pfacts(stream[j][1])
                        if not g["text"] and not g["math"]:
                            # allow blank lines only if the algorithm continues afterwards
                            k = j + 1
                            while k < n and stream[k][0] == "p" and not self._pfacts(stream[k][1])["text"]:
                                k += 1
                            if k < n and stream[k][0] == "p":
                                g2 = self._pfacts(stream[k][1])
                                if is_algo_line(g2["text"], g2["indented"], g2["numbered"], g2["math"], g2["style"], g2["mono"]) \
                                        and not is_algo_head(g2["text"], g2["bold"]):
                                    group += [stream[x][1] for x in range(j, k)]
                                    j = k
                                    continue
                            break
                        if is_algo_head(g["text"], g["bold"]) or CAPTION_RE.match(g["text"] or "~"):
                            break
                        if not is_algo_line(g["text"], g["indented"], g["numbered"], g["math"], g["style"], g["mono"]):
                            break
                        group.append(stream[j][1])
                        j += 1
                        if len(group) > 120:
                            break
                    if len(group) > 1 or f["breaks"] >= 2:
                        self._emit_algorithm(doc, group)
                        i = j
                        continue
            self._handle_para(Paragraph(el, self.d._body), clf, doc, body_size)
            i += 1
        return doc

    # ---------------- tables ----------------
    def _cell_paras(self, tbl_el):
        return [p for p in tbl_el.iter(qn("w:p"))]

    def _table_kind(self, el) -> str:
        paras = self._cell_paras(el)
        texts = []
        for p in paras:
            f = self._pfacts(p)
            if f["text"]:
                texts.append(f)
        # algorithm in a table
        for f in texts[:2]:
            if is_algo_head(f["text"], f["bold"]):
                return "algo_tbl"
        if texts and all(f["mono"] for f in texts) and len(texts) >= 3:
            return "algo_tbl"
        rows = el.findall(qn("w:tr"))
        cells = [tc for tr in rows for tc in tr.findall(qn("w:tc"))]
        # figure layout table: pictures + their caption / sub-figure labels → unwrap
        if not has_math(el) and (next(el.iter(qn("w:drawing")), None) is not None or
                                 next(el.iter(f"{{{NS_V}}}imagedata"), None) is not None):
            from .classify import is_caption_text
            pic_cells = text_cells = 0
            long_text = False
            for tc in cells:
                has_pic = next(tc.iter(qn("w:drawing")), None) is not None or \
                    next(tc.iter(f"{{{NS_V}}}imagedata"), None) is not None
                t = re.sub(r"\s+", " ", "".join(x.text or "" for x in tc.iter(qn("w:t")))).strip()
                if has_pic:
                    pic_cells += 1
                elif t and not is_caption_text(t):
                    text_cells += 1
                    long_text = long_text or len(t) > 80
            if pic_cells and text_cells <= pic_cells and not long_text:
                return "layout"
        # equation table: maths + only numbers / operators as text
        if has_math(el) or any(r.img and r.img.get("kind") == "ole-eq" for p in paras for r in self._runs(p, False)):
            ok = True
            for tc in cells:
                t = "".join(x.text or "" for x in tc.iter(qn("w:t"))).strip()
                if t and not (EQ_NUM_RE.match(t) or len(t) <= 3 or re.fullmatch(r"(?:where|and|s\.t\.|subject to)[,:]?", t, re.I)):
                    ok = False
                    break
            if ok:
                return "eq_tbl"
        if len(rows) == 1 and len(cells) <= 2:
            return "layout"
        return "tbl"

    def _emit_eq_table(self, doc, el):
        maths, num = [], None
        for p in self._cell_paras(el):
            for r in self._runs(p, False):
                if r.is_math:
                    maths.append(r)
                elif EQ_NUM_RE.match(r.text.strip() or "~"):
                    num = EQ_NUM_RE.match(r.text.strip()).group(1)
            t = "".join(x.text or "" for x in p.iter(qn("w:t"))).strip()
            if EQ_NUM_RE.match(t):
                num = EQ_NUM_RE.match(t).group(1)
        doc.blocks.append(M.Block(M.EQUATION, runs=maths, eq_num=num, src=el, raw=[copy.deepcopy(el)]))

    def _emit_algorithm(self, doc, elements):
        lines = self._algo_lines(elements)
        title = lines[0][1] if lines else []
        title_text = "".join(r.text for r in title if not r.is_math)
        if not is_algo_head(title_text, True):
            title = []
        else:
            lines = lines[1:]
        paras = [p for e in elements for p in ([e] if etree.QName(e).localname == "p" else e.iter(qn("w:p")))]
        def numbered(p):
            if p.find(qn("w:pPr") + "/" + qn("w:numPr")) is not None:
                return True
            ps = p.find(qn("w:pPr") + "/" + qn("w:pStyle"))
            return ps is not None and "number" in (ps.get(qn("w:val")) or "").lower()
        auto_num = sum(1 for p in paras if numbered(p))
        doc.blocks.append(M.Block(M.ALGORITHM, runs=title, lines=lines, raw=[copy.deepcopy(e) for e in elements],
                                  src=elements[0], orig=title_text, ordered=auto_num >= max(1, len(lines) // 2)))

    def _emit_table(self, doc, el):
        from docx.table import Table
        tbl = Table(el, self.d._body)
        rows, rows_runs = [], []
        complex_ = has_math(el) or el.find(".//" + qn("w:tbl")) is not None or \
            next(el.iter(qn("w:drawing")), None) is not None or \
            any(True for _ in el.iter(qn("w:gridSpan"))) or any(True for _ in el.iter(qn("w:vMerge")))
        for tr in el.findall(qn("w:tr")):
            cells, cell_runs = [], []
            for tc in tr.findall(qn("w:tc")):
                cr: List[M.Run] = []
                for k, p in enumerate(tc.findall(qn("w:p"))):
                    pr = self._runs(p, False)
                    if k and cr and pr:
                        cr.append(M.Run(" "))
                    cr += [r for r in pr if not (r.img and r.img.get("kind") == "pic")]
                cells.append(re.sub(r"\s+", " ", "".join(r.text for r in cr if not r.is_math)).strip())
                cell_runs.append(merge_adjacent(cr))
                span = tc.find(qn("w:tcPr") + "/" + qn("w:gridSpan")) if tc.find(qn("w:tcPr")) is not None else None
                if span is not None:
                    for _ in range(int(span.get(qn("w:val"), "1")) - 1):
                        cells.append(""); cell_runs.append(None)      # None = covered by a horizontal merge
            rows.append(cells)
            rows_runs.append(cell_runs)
        doc.blocks.append(M.Block(M.TABLE, rows=rows, rows_runs=rows_runs, src=el,
                                  raw=[copy.deepcopy(el)] if complex_ else None))

    # ---------------- paragraphs ----------------
    def _handle_para(self, p, clf: RoleClassifier, doc: M.Document, body_size: float):
        st = self._para_style(p)
        style_name = st.name if st is not None else ""
        style_bold = self._style_bold(st)
        runs = self._runs(p._p, style_bold)

        text = "".join(r.text for r in runs if not r.is_math).strip()
        objs = [r for r in runs if r.img is not None]
        only_num = (not text) or bool(EQ_NUM_RE.match(text))

        # decide what each embedded object is
        keep: List[M.Run] = []
        figures = []
        for r in runs:
            if r.img is None:
                keep.append(r)
                continue
            info = r.img
            small = info.get("h") is not None and info["h"] < SMALL_IMG_IN
            if info["kind"] == "ole-eq":
                keep.append(r)                     # MathType / Equation Editor: always maths
            elif small and info.get("inline", True) and (text or len(objs) == 1 and only_num and EQ_NUM_RE.match(text or "~")):
                keep.append(r)                     # inline symbol / equation picture
            elif info.get("blob") is not None:
                figures.append(r)
        for r in figures:
            info = r.img
            doc.blocks.append(M.Block(M.FIGURE, image=info["blob"], image_ext=info["ext"], width_in=info.get("w"),
                                      src=p._p))
        runs = keep
        maths = [r for r in runs if r.is_math]

        # display equation: only maths (+ optional "(n)" number) -> verbatim
        if maths and only_num:
            num = EQ_NUM_RE.match(text).group(1) if text else None
            doc.blocks.append(M.Block(M.EQUATION, runs=maths, eq_num=num, src=p._p,
                                      raw=None if figures else [copy.deepcopy(p._p)]))
            return
        if not text and not maths:
            return

        size = self._style_size(st)
        rsz = [r.font.size.pt for r in p.runs if r.font.size and r.text.strip()]
        if rsz:
            size = max(rsz)
        nonempty = [r for r in runs if (r.text.strip() or r.is_math)]
        bold = bool(nonempty) and all(r.bold for r in nonempty if not r.is_math)

        numpr = p._p.find(qn("w:pPr") + "/" + qn("w:numPr")) if p._p.find(qn("w:pPr")) is not None else None
        if numpr is None and st is not None:
            se = st.element.find(qn("w:pPr") + "/" + qn("w:numPr")) if st.element.find(qn("w:pPr")) is not None else None
            numpr = se if (style_name.lower().startswith("list")) else None
        is_list = numpr is not None or style_name.lower().startswith("list")
        ordered, lvl = False, 0
        if numpr is not None:
            nid = numpr.find(qn("w:numId"))
            il = numpr.find(qn("w:ilvl"))
            lvl = int(il.get(qn("w:val"))) if il is not None else 0
            if nid is not None:
                ordered = self._is_ordered(nid.get(qn("w:val")), str(lvl))
        if "number" in style_name.lower():
            ordered = True
        if style_name.lower().startswith("heading"):
            is_list = False

        info = ParaInfo(text=text, style=style_name, size=size, bold=bold, is_list=is_list,
                        list_ordered=ordered, list_level=lvl, is_title_style=style_name.lower() == "title")
        role, level, label, cleaned = clf.feed(info)
        if role is None:
            return
        if role == "abstract_label":
            doc.blocks.append(M.Block(M.ABSTRACT_HEADING, runs=runs, orig=text, src=p._p, src_size=size))
            return
        if role == M.HEADING and maths:
            role, level = M.PARA, 0              # a line with maths is never a section heading
        if cleaned is not None:
            runs = trim_leading(runs, cleaned)
            if not runs and role not in (M.KEYWORDS,):
                return
        runs = strip_trailing(runs)
        blk = M.Block(role, runs=runs, level=level, label=label, style_name=style_name, orig=text,
                      src=p._p, src_size=size, src_bold=bold)
        if role == M.LIST:
            blk.ordered = ordered
            blk.list_level = lvl
        if role == M.CAPTION:
            m = CAPTION_RE.match(text)
            kind = m.group("kind").lower()
            blk.cap_kind = "table" if kind.startswith("tab") else "figure"
            num = m.group("num")
            blk.cap_num = int(num) if num.isdigit() else None
            blk.runs = trim_leading(runs, m.group("rest")) if m.group("rest").strip() else []
        if role == M.HEADING:
            lvl_num, bare = strip_heading_number(text)
            if lvl_num and not style_name.lower().startswith("heading"):
                blk.level = lvl_num
            blk.runs = trim_leading(runs, bare)
        doc.blocks.append(blk)


def read_docx(path_or_stream) -> M.Document:
    return DocxReader(path_or_stream).read()
