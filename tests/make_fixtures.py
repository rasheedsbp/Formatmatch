"""Create realistic test manuscripts (APA-style Word draft) for end-to-end tests."""
import copy
import io
import os

import docx
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))


def fig_png():
    fig, ax = plt.subplots(figsize=(4, 2.5))
    ax.plot([1, 2, 3, 4], [0.71, 0.82, 0.88, 0.93], marker="o")
    ax.set_xlabel("Epoch"); ax.set_ylabel("Accuracy")
    b = io.BytesIO(); fig.savefig(b, format="png", dpi=120); plt.close(fig)
    b.seek(0)
    return b


def make_apa_draft(path):
    d = docx.Document()
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    p = d.add_paragraph(); r = p.add_run("A hybrid deep learning framework for intrusion detection in IoMT networks"); r.bold = True; r.font.size = Pt(16)
    d.add_paragraph("Ravi Kumar¹, Priya Sharma², and John Doe¹")
    d.add_paragraph("¹Department of Computer Science, Osmania University, Hyderabad, India")
    d.add_paragraph("²School of Engineering, University of Leeds, United Kingdom; ravi.k@ou.ac.in")
    d.add_paragraph("Abstract: The Internet of Medical Things (IoMT) is exposed to many attacks. "
                    "This study proposes a hybrid CNN–LSTM model that reaches 98.4% accuracy on CICIoMT2024. "
                    "Results show strong generalisation across attack families.")
    d.add_paragraph("Keywords: IoMT; intrusion detection; deep learning; CNN; LSTM")
    h = d.add_paragraph(); r = h.add_run("1. Introduction"); r.bold = True
    p = d.add_paragraph("Medical devices are increasingly connected (Smith et al., 2020). Prior work has used shallow "
                        "classifiers (Lee & Park, 2019; Brown, 2021), whereas Smith et al. (2020) used CNNs. ")
    r = p.add_run("Escherichia coli"); r.italic = True
    p.add_run(" is not relevant here but tests italics, and H")
    r = p.add_run("2"); r.font.subscript = True
    p.add_run("O tests subscripts.")
    d.add_paragraph("The contributions of this work are:")
    d.add_paragraph("A hybrid CNN–LSTM detector.", style="List Bullet")
    d.add_paragraph("An evaluation on two public datasets.", style="List Bullet")
    h = d.add_paragraph(); r = h.add_run("2. Methodology"); r.bold = True
    h = d.add_paragraph(); r = h.add_run("2.1 Dataset and preprocessing"); r.bold = True
    d.add_paragraph("The dataset statistics are summarised in Table 1 and the loss is defined in Equation (1).")
    # equation copied from pandoc docx
    eqd = docx.Document(os.path.join(HERE, "eq.docx"))
    eq_p = copy.deepcopy(eqd.paragraphs[0]._p)
    d.element.body.insert(len(d.element.body) - 1, eq_p)
    d.add_paragraph("Table 1: Dataset summary")
    t = d.add_table(rows=3, cols=3); t.style = "Table Grid"
    for i, row in enumerate([["Class", "Samples", "Share"], ["Benign", "120,000", "60%"], ["Attack", "80,000", "40%"]]):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    h = d.add_paragraph(); r = h.add_run("3. Results and Discussion"); r.bold = True
    d.add_paragraph("Figure 1 shows the training curve; accuracy rises steadily (Kumar, 2022).")
    d.add_picture(fig_png(), width=Inches(4.5))
    d.add_paragraph("Figure 1. Training accuracy over epochs.")
    h = d.add_paragraph(); r = h.add_run("4. Conclusion"); r.bold = True
    d.add_paragraph("The proposed model is accurate and lightweight (Lee & Park, 2019).")
    h = d.add_paragraph(); r = h.add_run("References"); r.bold = True
    for ref in [
        "Brown, T. (2021). Machine learning for network security. IEEE Access, 9, 1123–1135. https://doi.org/10.1109/ACCESS.2021.1234567",
        "Kumar, R. (2022). Lightweight intrusion detection for medical devices. Computers & Security, 115, 102611.",
        "Lee, J., & Park, S. (2019). Shallow classifiers for anomaly detection. Expert Systems with Applications, 120(3), 45–58. https://doi.org/10.1016/j.eswa.2019.01.010",
        "Smith, A. B., Jones, C., & Patel, D. (2020). Deep learning in the Internet of Medical Things. Journal of Network and Computer Applications, 150, 102480. https://doi.org/10.1016/j.jnca.2020.102480",
    ]:
        d.add_paragraph(ref)
    d.save(path)


def make_template(path):
    """A journal-style Word template: 1-col front matter, 2-col body, Arial headings, numbered refs."""
    from docx.enum.section import WD_SECTION
    d = docx.Document()
    st = d.styles["Normal"]; st.font.name = "Cambria"; st.font.size = Pt(9.5)
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.left_margin = sec.right_margin = Inches(0.7)
    def para(text, size=9.5, bold=False, italic=False, align=None, font=None, label=None, label_fmt=None, indent=None):
        p = d.add_paragraph()
        if align: p.alignment = align
        if indent is not None: p.paragraph_format.first_line_indent = Inches(indent)
        if label:
            r = p.add_run(label); r.font.size = Pt(size); r.bold, r.italic = label_fmt
            if font: r.font.name = font
        r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
        if font: r.font.name = font
        return p
    para("Template Title Goes Here", 20, True, align=WD_ALIGN_PARAGRAPH.LEFT, font="Arial")
    para("First Author, Second Author", 11, align=WD_ALIGN_PARAGRAPH.LEFT, font="Arial")
    para("Department of X, University of Y, City, Country", 8, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    para("This is the abstract text of the template paper, it should be about two hundred words long and describe the work.", 9,
         label="ABSTRACT: ", label_fmt=(True, False), align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    para("alpha; beta; gamma", 9, label="KEYWORDS: ", label_fmt=(True, False))
    # section break: rest is two columns
    new = d.add_section(WD_SECTION.CONTINUOUS)
    cols = new._sectPr.find(qn("w:cols"))
    if cols is None:
        from docx.oxml import OxmlElement
        cols = OxmlElement("w:cols"); new._sectPr.append(cols)
    cols.set(qn("w:num"), "2"); cols.set(qn("w:space"), "425")
    para("1 INTRODUCTION", 10, True, font="Arial")
    para("Body text of the template is set in Cambria 9.5 pt with a first line indent and justified alignment so that the "
         "renderer can learn it from this exemplar paragraph which is long enough to qualify as body text.", 9.5,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=0.15)
    para("1.1 Background", 9.5, True, italic=True, font="Arial")
    para("More body text appears here and cites a source [1] and another [2], which shows numbered citations in brackets.",
         9.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=0.15)
    para("Sample caption text.", 8, label="Figure 1 | ", label_fmt=(True, False), align=WD_ALIGN_PARAGRAPH.LEFT)
    para("Sample table caption.", 8, label="Table 1 | ", label_fmt=(True, False), align=WD_ALIGN_PARAGRAPH.LEFT)
    t = d.add_table(rows=2, cols=2); t.style = "Light Shading"
    para("REFERENCES", 10, True, font="Arial")
    para("[1] A. Author, “Title of paper,” Journal Name, vol. 1, no. 2, pp. 3–4, 2020.", 8)
    para("[2] B. Writer and C. Coder, “Another title,” in Proc. Conf., 2019, pp. 1–9.", 8)
    d.save(path)


if __name__ == "__main__":
    make_apa_draft(os.path.join(HERE, "draft_apa.docx"))
    make_template(os.path.join(HERE, "template_journal.docx"))
    print("ok")
